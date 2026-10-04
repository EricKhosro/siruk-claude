#!/usr/bin/env python3
"""Read-only price audit, 2026-10-03 (shop opens 2026-10-04).

Part A - register csv/AllAngineProduct-FINAL-2026-09-28.xlsx: every row -> live variant (same
  mapping as runs/2026-10-02-final/not-imported.py) -> live price vs `Վաճառքի գին` (col H);
  differs or blank -> hafo.am price for the article code. Also col I `Kg` vs the "1 kg" variant.
Part B - csv/Price Royal Canin Nor _ SIRUK.xlsx: rows with a number in column E -> live product
  (runs/2026-09-30-rc plan/state) -> pack price vs `Վաճառքի Գին`, 1 kg / single-pouch variant vs
  `Վաճառքի Գին կիլոգրամով`; blank price -> hafo.
Live prices: live.json (public storefront API = what a customer pays).
"""
import csv, json, os, re, subprocess, sys
from collections import Counter
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
live = json.load(open(os.path.join(HERE, "live.json"), encoding="utf-8"))["products"]
HAFO_CACHE = os.path.join(HERE, "hafo-cache.json")
hafo_cache = json.load(open(HAFO_CACHE, encoding="utf-8")) if os.path.exists(HAFO_CACHE) else {}
TIMES = "×"


def norm(s):
    return re.sub(r"\s+", " ", str(s or "").replace("\xa0", " ")).strip()


def num(x):
    if x is None or x == "":
        return None
    try:
        return float(str(x).replace(",", "").strip())
    except ValueError:
        return None


by_sku, by_vid = {}, {}
for pid, p in live.items():
    for v in p["variants"]:
        by_sku[norm(v["sku"])] = (pid, p, v)
        by_vid[v["id"]] = (pid, p, v)


def find(code):
    code = norm(code)
    if not code:
        return None
    for c in (code, re.sub(r"(Tx|TXN)$", "", code), code.lstrip("0"), code.replace(" ", ""),
              re.sub(r"^(MG|MM|GM|IN|BP|DL|CP|PG|PM|KF) ?", "", code), f"4{code}V", f"{code}V"):
        if c in by_sku:
            return by_sku[c]
    return None


def hafo(code):
    if not code:
        return None
    if code not in hafo_cache:
        r = subprocess.run([sys.executable, os.path.join(ROOT, "scripts/hafo-lookup.py"), "--code", code],
                           capture_output=True, text=True, encoding="utf-8", cwd=ROOT)
        try:
            d = json.loads(r.stdout)
        except Exception:
            d = {"error": (r.stderr or r.stdout)[-300:]}
        hafo_cache[code] = {k: d.get(k) for k in ("confirmed", "price_source", "price_amd", "wholesale_price_amd",
                                                  "variant_kg_price", "variant_name_hy", "variant_sku", "url",
                                                  "listing_price_amd", "error", "matched_by")}
        json.dump(hafo_cache, open(HAFO_CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=0)
    return hafo_cache[code]


def hafo_price(h):
    if not h or not h.get("confirmed"):
        return None, ("hafo: not found" if h else "hafo: no code")
    if h.get("price_source") != "variant":
        return None, f"hafo: no per-variant price (listing {h.get('listing_price_amd')})"
    return h.get("price_amd"), ""


def promo(v):
    bits = []
    if v.get("salePrice") not in (None, v.get("price")):
        bits.append(f"salePrice {v['salePrice']}")
    if v.get("discount"):
        bits.append(f"discount {v['discount']}")
    if v.get("promotion"):
        bits.append("promotion")
    if v.get("compareAtPrice"):
        bits.append(f"compareAt {v['compareAtPrice']}")
    return "; ".join(bits)


def i(x):
    return "" if x is None else int(round(x))


# ================================================================= Part A: register
wb = openpyxl.load_workbook(os.path.join(ROOT, "csv/AllAngineProduct-FINAL-2026-09-28.xlsx"), data_only=True)
rows = list(wb.worksheets[0].iter_rows(values_only=True))
hdr_i = next(k for k, r in enumerate(rows) if r and r[0] == "Կոդ")
reg = [r for r in rows[hdr_i + 1:] if r and r[0] and r[1]]
status = {x["reg_no"]: x for x in json.load(open(os.path.join(ROOT, "state/register/register-status.json"), encoding="utf-8"))}
over = json.load(open(os.path.join(ROOT, "state/register/match-overrides.json"), encoding="utf-8"))
found = {x["Կոդ"]: x for x in csv.DictReader(open(os.path.join(ROOT, "runs/2026-10-01-register/found.csv"), encoding="utf-8-sig"))}

A, A_kg, A_notlive = [], [], []
seen_vid = {}
for r in reg:
    no, name, cost, sale, kg = str(r[0]), norm(r[1]), num(r[5]), num(r[7]), num(r[8])
    code = (found.get(no) or {}).get("Article code") or over.get(no) or (status.get(no) or {}).get("code")
    code = code or {"00030": "REG00030"}.get(no)
    if not code:
        m = re.findall(r"/\s*(\d{4,8})\s*$", name)
        code = m[0] if m and find(m[0]) else None
    hit = find(code) if code else None
    st = status.get(no) or {}
    if not hit and st.get("live") and st.get("variant_id") in by_vid:
        hit = by_vid[st["variant_id"]]
    if not hit:
        A_notlive.append({"Կոդ": no, "Անվանում": name, "article": code or "", "cost": i(cost), "Վաճառքի գին": i(sale)})
        continue
    pid, p, v = hit
    seen_vid.setdefault(v["id"], []).append(no)
    lp = v["price"]
    base = {"Կոդ": no, "Անվանում": name, "article": code or "", "product_id": pid,
            "product": f"{p['brand']} {p['name']}", "variant": v["label"], "sku": v["sku"],
            "cost": i(cost), "Վաճառքի գին": i(sale), "live price": lp, "promo": promo(v)}
    if sale is not None and lp == sale:
        A.append(dict(base, hafo="", verdict="OK" if not base["promo"] else "OK (but promo on live)", note=""))
    else:
        hp, hnote = hafo_price(hafo(code))
        if sale is None:
            if hp is None:
                verdict, note = "NO REGISTER PRICE, NO HAFO PRICE", hnote
            elif hp == lp:
                verdict, note = "OK vs hafo (register blank)", ""
            else:
                verdict, note = "DIFF vs hafo (register blank)", f"live - hafo = {lp - hp:+}"
        else:
            note = f"live - register = {lp - sale:+.0f}"
            if hp is not None:
                note += "; hafo " + ("= register" if hp == sale else ("= live" if hp == lp else "matches neither"))
            else:
                note += f"; {hnote}"
            verdict = "DIFF vs register"
        if cost is not None and lp <= cost:
            note += "; LIVE PRICE <= COST"
        A.append(dict(base, hafo=i(hp), verdict=verdict, note=note))
    # ---- per-kg column vs the "1 kg" variant
    kgv = by_sku.get(norm(v["sku"]) + "-KG")
    if kgv is None:
        kgv = next(((pid, p, x) for x in p["variants"] if x is not v and norm(x["sku"]).startswith(norm(v["sku"]))
                    and x.get("size") == "1 kg"), None)
    if kg is not None or kgv is not None:
        kp = kgv[2]["price"] if kgv else None
        if kg is not None and kp == kg:
            vd = "OK"
        elif kg is None:
            vd = "1 kg variant live but register Kg blank"
        elif kgv is None:
            vd = "register has Kg price but no 1 kg variant live"
        else:
            vd = "DIFF"
        A_kg.append({"Կոդ": no, "Անվանում": name, "product_id": pid, "bag sku": v["sku"], "bag": v["label"],
                     "register Kg": i(kg), "1 kg sku": kgv[2]["sku"] if kgv else "",
                     "1 kg live price": kp if kgv else "",
                     "hafo kg_price": i((hafo(code) or {}).get("variant_kg_price")) if vd != "OK" else "",
                     "verdict": vd})

dupes = {vid: nos for vid, nos in seen_vid.items() if len(nos) > 1}

# ================================================================= Part B: Royal Canin
ws = openpyxl.load_workbook(os.path.join(ROOT, "csv/Price Royal Canin Nor _ SIRUK.xlsx"), data_only=True).worksheets[0]
rc = list(ws.iter_rows(values_only=True))
plan = json.load(open(os.path.join(ROOT, "runs/2026-09-30-rc/plan.json"), encoding="utf-8"))["products"]
rcstate = json.load(open(os.path.join(ROOT, "runs/2026-09-30-rc/state.json"), encoding="utf-8"))["done"]
row2 = {}
for pr in plan:
    for v in pr["variants"]:
        row2.setdefault(v["_row"], {"key": pr["key"], "skus": []})["skus"].append(v["sku"])
B = []
for n, r in enumerate(rc[1:], 2):
    E = r[4]
    if not isinstance(E, (int, float)):
        continue
    name, cost, sale, perkg = norm(r[0]), num(r[1]), num(r[2]), num(r[3])
    rec = {"row": n, "name": name, "E": E, "cost": i(cost), "Վաճառքի Գին": i(sale), "Վաճառքի Գին կգ": i(perkg)}
    m = row2.get(n)
    pid = str((rcstate.get(m["key"]) or {}).get("product_id")) if m else None
    p = live.get(pid) if pid else None
    if not p:
        for s in (m or {}).get("skus", []):
            if norm(s) in by_sku:
                pid, p, _ = by_sku[norm(s)]
                break
    if not p:
        B.append(dict(rec, product="", checks="NOT LIVE on siruk.am" + ("" if m else " (row was not in the 2026-09-30 import plan)")))
        continue
    rec["product"] = f"{pid} {p['name']}"
    rec["live variants"] = " | ".join(f"{x['label']} [{x['sku']}] {x['price']} (stock {x['stock']})" for x in p["variants"])
    vs = p["variants"]
    roots = {s.split("-")[0] for s in (m["skus"] if m else [])}
    mine = [x for x in vs if x["sku"].split("-")[0] in roots] or vs
    checks = []
    if E == 1:
        bag = [x for x in mine if x.get("size") != "1 kg" and not x["sku"].endswith("-KG")]
        one = [x for x in mine if x.get("size") == "1 kg" or x["sku"].endswith("-KG")]
        exp, src = sale, "sheet"
        if exp is None:
            exp, _ = hafo_price(hafo(sorted(roots)[0] if roots else ""))
            src = "hafo"
        for x in bag:
            ok = exp is not None and x["price"] == exp
            checks.append(f"bag {x['label']}: live {x['price']} vs {src} {i(exp)} -> {'OK' if ok else 'DIFF'}")
            if cost is not None and x["price"] <= cost:
                checks.append(f"bag {x['label']} price <= cost")
        if not bag:
            checks.append("bag variant: MISSING")
        if perkg is not None:
            for x in one:
                checks.append(f"1 kg: live {x['price']} vs sheet {i(perkg)} -> {'OK' if x['price'] == perkg else 'DIFF'}")
            if not one:
                checks.append(f"1 kg variant MISSING (sheet per-kg {i(perkg)})")
        elif one:
            checks.append(f"1 kg variant live at {one[0]['price']} but sheet per-kg blank")
    else:  # cases of 12, sold one by one (runs/2026-10-01-rc/single-pouch.py)
        single = [x for x in mine if TIMES not in (x.get("size") or "") and TIMES not in (x["label"] or "")]
        case = [x for x in mine if x not in single]
        exp = perkg if perkg is not None else sale   # D = per piece; Mother & Babycat: C is per can
        for x in single:
            checks.append(f"single {x['label']}: live {x['price']} vs sheet {i(exp)} -> {'OK' if x['price'] == exp else 'DIFF'}")
        if not single:
            checks.append("single-piece variant MISSING")
        for x in case:
            orderable = (x.get("stock") or 0) > 0
            checks.append(f"case {x['label']} still on storefront at {x['price']}, stock {x['stock']}"
                          + (" -> ORDERABLE" if orderable else " (stock 0)"))
    for x in mine:
        if promo(x):
            checks.append(f"promo on {x['label']}: {promo(x)}")
    rec["checks"] = "; ".join(checks)
    B.append(rec)


def write(fn, rows):
    if not rows:
        return
    keys = list(dict.fromkeys(k for r in rows for k in r))
    with open(os.path.join(HERE, fn), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=keys); w.writeheader(); w.writerows(rows)


write("A-register-prices.csv", A)
write("A-register-1kg.csv", A_kg)
write("A-register-not-live.csv", A_notlive)
write("B-royal-canin.csv", B)
json.dump({"dupes": dupes}, open(os.path.join(HERE, "A-dupes.json"), "w"), ensure_ascii=False)
print("A:", len(reg), "rows;", len(A), "live;", len(A_notlive), "not live;", Counter(x["verdict"] for x in A))
print("A kg:", Counter(x["verdict"] for x in A_kg))
print("A rows sharing one variant:", len(dupes))
print("B:", len(B), "rows")
