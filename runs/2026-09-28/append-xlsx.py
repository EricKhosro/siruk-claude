"""Append the 29 sirook.pdf invoice rows to csv/AllAngineProduct-FINAL-2026-09-28.xlsx.

Rows go above the `Ընդամենը` totals row, numbered on from the last register
code, styled like the row above them; the totals are bumped by the new qty and
amount. Run with the scratch venv's python (needs openpyxl):

    <venv>/bin/python runs/2026-09-28/append-xlsx.py
"""
import copy, csv, json, os, runpy, subprocess

import openpyxl

RUN = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(RUN))
XLSX = os.path.join(ROOT, "csv/AllAngineProduct-FINAL-2026-09-28.xlsx")
PLAN = runpy.run_path(os.path.join(RUN, "plan.py"))["PRODUCTS"]
STATE = json.load(open(os.path.join(RUN, "state.json")))["done"]
HAFO = {r["csv"]["Article Code"]: r for r in json.load(open(os.path.join(RUN, "hafo.json")))}
ROWS = list(csv.DictReader(open(os.path.join(RUN, "sirook.csv"))))

def variant_ids(pid):
    out = subprocess.run([os.path.join(ROOT, "scripts/api.sh"), "GET", f"/products/{pid}"],
                         capture_output=True, text=True, cwd=ROOT).stdout
    return {v["sku"]: v["id"] for v in json.loads(out[out.find("{"):])["data"]["variants"]}


by_code = {}
for p in PLAN:
    for v in p["variants"]:
        by_code[v["code"]] = (p, v)

wb = openpyxl.load_workbook(XLSX)
ws = wb.active
total_row = ws.max_row
assert ws.cell(total_row, 1).value == "Ընդամենը", "totals row moved"
tmpl = total_row - 1
last_code = int(ws.cell(tmpl, 1).value)
ws.insert_rows(total_row, amount=len(ROWS))

qty_sum = amt_sum = 0
for i, r in enumerate(ROWS):
    row = total_row + i
    for col in range(1, 16):
        src, dst = ws.cell(tmpl, col), ws.cell(row, col)
        dst._style = copy.copy(src._style)
    code, cost, qty = r["Article Code"], int(r["Buy Price (AMD)"]), int(r["Qty Received"])
    qty_sum += qty; amt_sum += cost * qty
    vals = [f"{last_code + 1 + i:05d}", r["Product Name (as printed)"], "001", "հատ",
            qty, cost, cost * qty]
    if code in by_code:
        p, v = by_code[code]
        pid = STATE[p["key"]]["id"]
        vid = variant_ids(pid)[v.get("sku", v["code"])]
        link = f"https://demo.siruk.am/product/{STATE[p['key']]['slug']}/dp/{vid}/"
        brand = "Versele-Laga" if p["brand"] == 36 else "Trixie"
        note = (f"From the supplier invoice sirook.pdf (2026-09-28); hafo-confirmed EAN for {code}, "
                + ("same EAN on the official versele.com page" if brand == "Versele-Laga" and "versele.com" in p["source"]
                   else "old pack generation (versele.com only lists the 2025 relaunch EANs)" if brand == "Versele-Laga"
                   else "GTIN field on the official trixie.de page"))
        vals += [v["price"], None, "Yes", "", f"{p['name']['en']} — {v['label']}", link, v["ean"], note]
    else:  # 6004Tx: identified, no sale price anywhere
        vals += [None, None, "No",
                 "No sale price: not on hafo; zoovet.am and nemo.am don't carry it; no priced sibling → needed: the sale price for Trixie Lick made of Himalaya Salt 60 g (6004)",
                 "Lick made of Himalaya Salt — 60 g", None, "4011905060040",
                 "GTIN from the official trixie.de page for item no. 6004 (also carrefour.es, dammers.com)"]
    for col, val in enumerate(vals, 1):
        ws.cell(row, col).value = val
    if vals[12]:
        ws.cell(row, 13).hyperlink = vals[12]

tot = ws.max_row
assert ws.cell(tot, 1).value == "Ընդամենը"
ws.cell(tot, 5).value = round((ws.cell(tot, 5).value or 0) + qty_sum, 3)
ws.cell(tot, 7).value = (ws.cell(tot, 7).value or 0) + amt_sum
wb.save(XLSX)
print(f"appended {len(ROWS)} rows ({last_code + 1:05d}–{last_code + len(ROWS):05d}); totals +{qty_sum} qty, +{amt_sum} AMD")
