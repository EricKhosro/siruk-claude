import json, re, copy, collections, os, openpyxl
from openpyxl.styles import Font
S='runs/2026-09-25/'
C='.siruk-cache/'
status=json.load(open('state/register/register-status.json'))
snap=json.load(open(C+'catalogue-snapshot.json'))
var_by_id={}; var_by_sku=collections.defaultdict(list)
for p in snap.values():
    for v in p['variants']:
        var_by_id[v['id']]=(p,v)
        if v.get('sku'): var_by_sku[v['sku'].strip()].append((p,v))

def ean_ok(b):
    b=re.sub(r'\D','',str(b or ''))
    if len(b) not in (8,12,13,14): return None
    d=[int(x) for x in b]; s=sum(x*(3 if i%2==0 else 1) for i,x in enumerate(reversed(d[:-1])))
    return b if (10-s%10)%10==d[-1] else None
def norm(code):
    c=(code or '').strip().upper()
    m=re.match(r'^TX\s*(\d{6})$',c)          # hafo's Trixie sku: 5-digit article + 1 variant digit
    if m: return 'T:'+(m.group(1)[:5].lstrip('0'))
    c=c.replace(' ','')
    m=re.match(r'^(\d+)(TXN|TX)$',c)
    if m: return 'T:'+m.group(1).lstrip('0')
    return c
src=collections.defaultdict(set)
def add(code,b,s):
    b=ean_ok(b)
    if b and code: src[norm(code)].add((b,s))
def add_hafo_item(key,v):
    # only a listing hafo confirmed for our code, and only the barcode whose own sku IS that code
    if not v.get('confirmed'): return
    sk,bc=v.get('skus') or [],v.get('barcodes') or []
    if len(sk)!=len(bc): return
    arts=v.get('articles') or []
    for i,(s_,b) in enumerate(zip(sk,bc)):
        add(s_,b,'hafo')
        m=re.match(r'^([A-Z]{2})[\s\xa0]+(\d+)$',s_.strip())
        if m and m.group(1)!='TX': add(m.group(1)+m.group(2),b,'hafo'); add(m.group(2),b,'hafo')
        if len(arts)==len(sk) and arts[i]: add(arts[i],b,'hafo')
for k,v in json.load(open(C+'hafo-all.json')).items(): add_hafo_item(k,v)
if os.path.exists(S+'hafo-live.jsonl'):
    for line in open(S+'hafo-live.jsonl'):
        try: j=json.loads(line)
        except Exception: continue
        add_hafo_item(j['code'],j['r'])
for f in ('hafo-variant-barcode.json','gap-barcodes.json'):
    for k,v in json.load(open(C+f)).items(): add(k,v.get('barcode'),'hafo')
for k,v in json.load(open(C+'ean-images.json')).items(): add(k,v.get('ean'),'retailer page')
pinned={}
for r in status:
    h=r.get('hafo') or {}
    if h.get('accepted') and h.get('barcode') and r.get('code') and norm(h.get('sku'))==norm(r['code']):
        add(r['code'],h['barcode'],'hafo'); pinned[r['reg_no']]=ean_ok(h['barcode'])
rows={r['reg_no']:r for r in status}
import csv
holds=collections.defaultdict(list)
for o in csv.DictReader(open('state/open-items.csv')):
    if o['Area']=='import: hold': holds[o['Article code']].append(o['Issue']+': '+o['Detail'])
name_cands=json.load(open('state/register/hafo-name-candidates.json'))
def fmt(n): return f"{n:,.0f}"
pm={}
for o in csv.DictReader(open('state/pm-problem-list-2026-09-23.csv',encoding='utf-8-sig')):
    if o['Problem'].startswith('Not on the site') and o['Register row']:
        pm[o['Register row']]=f"{o['Problem'].replace('Not on the site — ','').capitalize()}: {o['Details'].strip()} → needed: {o['What we need from you'].strip()}"
def why_not(r, code):
    if code and holds.get(code): return 'On hold — ' + ' | '.join(holds[code])
    if r['reg_no'] in pm: return pm[r['reg_no']]
    if code and holds.get(code):
        return 'On hold — ' + ' | '.join(holds[code])
    if code:
        cost, sale = r.get('cost') or 0, r.get('sale_price') or 0
        if sale and cost and sale <= cost:
            return f'Sale price in the register ({fmt(sale)}) is not above our cost ({fmt(cost)}) — needs the correct sale price'
        if not sale:
            return 'No sale price in the register'
        return f'Identified (article {code}) with a valid price — not imported yet, in the import queue'
    msg = 'Could not identify the exact product: the register has no article code and no hafo / zoovet / nemo listing matches brand, flavour and size'
    cs = (name_cands.get(r["reg_no"]) or {}).get('candidates') or []
    if cs:
        b = cs[0]; msg += f'. Closest hafo listing: "{b["name"]}" ({"same cost" if b.get("cost_eq") else "different cost"}) — not the same product'
    return msg
wb=openpyxl.load_workbook('/Users/conceptmacmini01/Downloads/AllAngineProduct.xlsx')
ws=wb.active
hdr=ws['I3']; body=ws['I4']
for i,h in enumerate(['In Siruk?','Why not imported','English name (Siruk)','Link on demo.siruk.am','Barcode (EAN)']):
    ws.cell(3,10+i,h)._style=copy.copy(hdr._style)
for col,w in zip('JKLMN',(10,70,55,70,17)): ws.column_dimensions[col].width=w
stats=collections.Counter(); problems=[]; bc_rows=collections.defaultdict(set)
for rr in range(4,ws.max_row+1):
    reg=ws.cell(rr,1).value
    if reg is None or not re.match(r'^\d+$',str(reg)): continue
    r=rows.get(str(reg))
    if not r or r['name_hy'].strip()!=str(ws.cell(rr,2).value).strip():
        problems.append((reg,'register row mismatch')); continue
    code=r.get('code'); pv=None
    if r.get('variant_id') in var_by_id: pv=var_by_id[r['variant_id']]
    elif code and len(var_by_sku.get(code.strip(),[]))==1: pv=var_by_sku[code.strip()][0]
    name=link=None
    if pv:
        p,v=pv; lab=(v.get('label') or '').strip()
        name=p['name']+(f' — {lab}' if lab and lab.lower() not in p['name'].lower() else '')
        link=f"https://demo.siruk.am/product/{p['slug']}/dp/{v['id']}/"
        stats['linked']+=1
        if not code and v.get('sku'): code=v['sku']
    elif r.get('live'): problems.append((reg,'status says live, variant not in catalogue'))
    insiruk='Yes' if pv else 'No'
    why=None if pv else (why_not(r,code) if not r.get('live') else 'Was on the site on 2026-09-23 but is not in the live catalogue now — check')
    bc=None
    if code:
        cands={b for b,_ in src.get(norm(code),())}
        if not cands and re.fullmatch(r'\d{4,5}',code.strip()):   # Trixie code written without 'Tx': the EAN must embed it
            cands={b for b,_ in src.get('T:'+code.strip().lstrip('0'),()) if b[7:12]==code.strip().zfill(5)}
        if len(cands)==1: bc=next(iter(cands))
        elif len(cands)>1:
            if pinned.get(r['reg_no']) in cands: bc=pinned[r['reg_no']]
            else: problems.append((reg,f'{code}: conflicting barcodes {sorted(cands)}'))
    if bc: stats['barcode']+=1; bc_rows[bc].add(norm(code))
    for i,val in enumerate((insiruk,why,name,link,bc)):
        c=ws.cell(rr,10+i,val); c._style=copy.copy(body._style)
        if i==4 and val: c.number_format='@'
        if i==3 and val:
            c.hyperlink=val; c.font=Font(name=body.font.name,sz=body.font.sz,color='0563C1',underline='single')
    stats['rows']+=1; stats['in:'+insiruk]+=1
    if why: stats['why:'+why.split(' — ')[0].split(':')[0][:40]]+=1
    stats['has code' if code else 'no code']+=1
dup={b for b,cs in bc_rows.items() if len(cs)>1}
for rr in range(4,ws.max_row+1):
    if ws.cell(rr,14).value in dup:
        problems.append((ws.cell(rr,1).value,'barcode shared with another article: '+ws.cell(rr,14).value)); ws.cell(rr,14).value=None; stats['barcode']-=1
ws.freeze_panes='A4'
out='/Users/conceptmacmini01/Downloads/AllAngineProduct-FINAL-2026-09-25.xlsx'
wb.save(out); print(stats); json.dump(problems,open(S+'problems.json','w'),ensure_ascii=False,indent=0)
print(collections.Counter(p[1].split(':')[0] if 'conflict' not in p[1] else 'conflict' for p in problems))
print([p for p in problems if 'conflict' in p[1]])
