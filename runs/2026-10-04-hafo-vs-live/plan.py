#!/usr/bin/env python3
"""Target price for every live siruk.am variant (user, 2026-10-04):
  1. matches hafo.am (article code)        -> hafo's regular price (1 kg variant -> hafo kg_price)
  2. else Royal Canin sheet row (col E set) -> its `Վաճառքի Գին` (1 kg / single pouch -> `... կիլոգրամով`)
  3. else register xlsx `Վաճառքի գին`      -> that price (1 kg variant -> `Kg` column)
  4. else                                   -> unchanged, listed
Only csv/AllAngineProduct-FINAL-2026-09-28.xlsx and csv/Price Royal Canin Nor _ SIRUK.xlsx are used
(no Product.numbers, no zoovet/nemo). Writes nothing. -> plan.csv, plan.json"""
import csv, json, os, re
from collections import Counter
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
live = json.load(open(os.path.join(HERE, "live.json"), encoding="utf-8"))["products"]
H = {r["sku"]: r for r in csv.DictReader(open(os.path.join(HERE, "all-variants-vs-hafo.csv"), encoding="utf-8-sig"))}
urls = json.load(open(os.path.join(HERE, "siruk-urls.json"), encoding="utf-8"))
FALSE_HAFO = {"12767", "25187", "25252", "4161", "25251", "23472"}   # Trixie: hafo row is another product

# ---- register xlsx: sku -> row (mapping from the 2026-10-03 audit, sku = article code on all 1,354)
ws = openpyxl.load_workbook(os.path.join(ROOT, "csv/AllAngineProduct-FINAL-2026-09-28.xlsx"), data_only=True).worksheets[0]
xl = {str(r[0]): r for r in ws.iter_rows(values_only=True) if r and r[0] and r[1]}
reg_by_sku = {}
for r in csv.DictReader(open(os.path.join(ROOT, "runs/2026-10-03-price-audit/A-register-prices.csv"), encoding="utf-8-sig")):
    reg_by_sku.setdefault(r["sku"], []).append(r["Կոդ"])

# ---- Royal Canin sheet: sku -> (row, C, D, E) via the 2026-09-30 import plan
rcws = openpyxl.load_workbook(os.path.join(ROOT, "csv/Price Royal Canin Nor _ SIRUK.xlsx"), data_only=True).worksheets[0]
rcrows = {n: r for n, r in enumerate(rcws.iter_rows(values_only=True), 1)}
plan_rc = json.load(open(os.path.join(ROOT, "runs/2026-09-30-rc/plan.json"), encoding="utf-8"))["products"]
rc_by_sku = {}
for pr in plan_rc:
    for v in pr["variants"]:
        rc_by_sku[v["sku"]] = v["_row"]
# single pouches that got their own EAN sku on 2026-10-01 (runs/2026-10-01-rc/single-pouch-state.json)
sp = json.load(open(os.path.join(ROOT, "runs/2026-10-01-rc/single-pouch-state.json"), encoding="utf-8"))["done"]
for pid, d in sp.items():
    for line in d.get("plan", []):
        m = re.match(r"retire (\S+) .* single (\S+) ", line)
        if m and m.group(1) in rc_by_sku:
            rc_by_sku.setdefault(m.group(2), rc_by_sku[m.group(1)])


def num(x):
    try:
        return float(str(x).replace(",", "")) if x not in (None, "") else None
    except ValueError:
        return None


rows = []
for pid, p in live.items():
    for v in p["variants"]:
        sku, lp = str(v["sku"]).strip(), v["price"]
        kind = "1 kg" if re.search(r"-(KG|1KG)$", sku) else ("single" if sku.endswith("-1") else "pack")
        base = re.sub(r"-(KG|1KG)$", "", sku)
        h = H.get(sku, {})
        regs = reg_by_sku.get(base, [])
        x = xl.get(regs[0]) if regs else None
        target = src = None
        note = ""
        # 1. hafo
        if h.get("hafo price") and sku not in FALSE_HAFO and h["group"] != "no hafo row":
            target, src = int(float(h["hafo price"])), "hafo.am"
            xc = num(x[5]) if x else None
            hw = num(h.get("hafo wholesale"))
            if kind != "1 kg" and xc is not None and hw is not None and abs(xc - hw) >= 1:
                note = f"xlsx cost {int(xc)} != hafo wholesale {int(hw)}"
        # 2. Royal Canin sheet
        elif sku in rc_by_sku and isinstance(rcrows[rc_by_sku[sku]][4], (int, float)):
            n = rc_by_sku[sku]; r = rcrows[n]
            C, D = num(r[2]), num(r[3])
            if kind == "1 kg" or kind == "single" or (D is not None and "×" not in v["label"] and v.get("size") in ("85 g",)):
                target = D
            else:
                target = C
            if r[4] == 12 and "×" in (v["label"] or ""):
                target, note = C, "retired 12-pack case (stock 0)"
            src = f"Royal Canin sheet row {n}"
            if target is None:
                src, note = None, f"Royal Canin sheet row {n} has no price for this variant"
        # 3. register xlsx
        elif x is not None:
            t = num(x[8]) if kind == "1 kg" else num(x[7])
            if t is not None:
                target, src = t, f"register xlsx row {regs[0]} ({'Kg' if kind == '1 kg' else 'Վաճառքի գին'})"
            else:
                note = f"register xlsx row {regs[0]} has no {'Kg' if kind == '1 kg' else 'Վաճառքի գին'}"
        else:
            note = note or "not on hafo, not in either file"
        target = int(target) if target is not None else None
        cost = num(h.get("our cost"))
        flags = []
        if target is not None and target % 10:
            flags.append("not a multiple of 10")
        if target is not None and cost is not None and target <= cost:
            flags.append(f"target <= cost {int(cost)}")
        action = ("unchanged (no source)" if target is None else ("OK" if target == lp else "CHANGE"))
        u = urls[pid]
        rows.append({"action": action, "product_id": pid, "variant_id": v["id"], "sku": sku, "kind": kind,
                     "Product name": f"{p['brand']} {p['name']} — {v['label']}",
                     "Link in Siruk": "https://siruk.am" + (u["by_sku"].get(sku) or u["product_url"]),
                     "current price": lp, "new price": target if target is not None else "",
                     "change": (target - lp) if target is not None else "", "source": src or "",
                     "hafo link": h.get("hafo url", "") if src == "hafo.am" else "",
                     "register row": " ".join(regs), "our cost": h.get("our cost", ""),
                     "stock": v.get("stock"), "note": "; ".join(filter(None, [note] + flags))})

rows.sort(key=lambda r: (r["action"] != "CHANGE", r["source"], r["Product name"]))
with open(os.path.join(HERE, "plan.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
json.dump([r for r in rows if r["action"] == "CHANGE"], open(os.path.join(HERE, "plan.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=0)
print(Counter(r["action"] for r in rows))
print("changes by source:", Counter(r["source"].split(" row")[0] for r in rows if r["action"] == "CHANGE"))
print("flags:", [(r["sku"], r["note"]) for r in rows if "multiple" in r["note"] or "<= cost" in r["note"]])
print("cost notes:", sum("xlsx cost" in r["note"] for r in rows))
