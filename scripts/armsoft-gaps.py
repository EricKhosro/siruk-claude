#!/usr/bin/env python3
"""List ArmSoft goods missing a barcode or a sale price, per unit.

    uv run --with openpyxl python scripts/armsoft-gaps.py <goods.xml> <prices.xlsx> [--out file.xlsx]

<goods.xml>   ArmSoft goods export (list "Ապրանքներ և ծառայություններ", select
              all rows with Ctrl+A, ribbon Արտահանել → Արտահանել XML).
<prices.xlsx> ArmSoft price list (Վաճառքներ → Գներ, period from 01/01/01,
              Արտահանել Excel → առանց ձևավորման).

A good with two units (bag 003 + kg 303, piece 001 + case 002 …) needs a
barcode AND a sale price (price type 01) on each unit. Prints a summary and,
with --out, writes one row per missing good + unit.

Also checks the ArmSoft hard rules (armsoft-sync skill -> "Hard rules"): Name
under 50 characters, VAT off, and a good with a kg unit (303) has it as main
unit, IsWeight on, and alt unit 003 whose coefficient = the bag kg in
Scale/scale-codes-per-kg.csv. With --out those go to a second sheet "Rules".
"""
import csv
import html
import re
import sys
from collections import defaultdict

import openpyxl

UNIT_NAMES = {'001': 'piece', '002': 'case', '003': 'bag', '303': 'kg'}
NAME_MAX = 50          # HDM (fiscal receipt) limit
BAG_KG = 'Scale/scale-codes-per-kg.csv'


def text(block, tag):
    m = re.search(rf'<{tag}>(.*?)</{tag}>', block, re.S)
    return html.unescape(m.group(1).strip()) if m else None


def parse_goods(path):
    xml = open(path, encoding='utf-8-sig').read()
    goods = []
    for g in re.findall(r'<Goods[ >].*?</Goods>', xml, re.S):
        head = g.split('<GoodQuantityUnits', 1)[0]
        units = [text(p, 'Unit') for p in
                 re.findall(r'<QuantityUnitPart>.*?</QuantityUnitPart>', g, re.S)]
        barcodes = defaultdict(list)
        for b in re.findall(r'<MTBarCode>.*?</MTBarCode>', g, re.S):
            if text(b, 'BarCode'):
                barcodes[text(b, 'Unit')].append(text(b, 'BarCode'))
        coefs = {text(p, 'Unit'): float(text(p, 'Coef') or 0) for p in
                 re.findall(r'<QuantityUnitPart>.*?</QuantityUnitPart>', g, re.S)}
        main_unit = text(head, 'Unit')
        if not units:
            units = [main_unit]
        goods.append({
            'code': text(head, 'Code'), 'name': text(head, 'Name'),
            'long': text(head, 'LongName'), 'type': text(head, 'Type'),
            'unit': main_unit, 'units': units, 'main_barcode': text(head, 'BarCode'),
            'is_weight': text(head, 'IsWeight') == 'true', 'barcodes': dict(barcodes),
            'vat': text(head, 'VAT') == 'true', 'alt_unit': text(head, 'AltUnit'),
            'coefs': coefs,
        })
    return goods


def parse_prices(path):
    """Latest sale price (type 01) per (code, unit)."""
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if r and 'Կոդ' in r)
    h = rows[hdr]
    col = {name: h.index(name) for name in ('Ամսաթիվ', 'Կոդ', 'Գնի տեսակ', 'Միավոր', 'Գին')}
    latest = {}
    for r in rows[hdr + 1:]:
        if not r or r[col['Կոդ']] is None:
            continue
        if str(r[col['Գնի տեսակ']]) != '01':
            continue
        key = (str(r[col['Կոդ']]), str(r[col['Միավոր']]))
        day = r[col['Ամսաթիվ']]
        if key not in latest or (day and latest[key][0] and day >= latest[key][0]):
            latest[key] = (day, r[col['Գին']])
    return {k: v[1] for k, v in latest.items()}


def bag_kg():
    """good code -> bag kg, for the goods Siruk sells loose."""
    try:
        rows = csv.DictReader(open(BAG_KG, encoding='utf-8-sig'))
    except FileNotFoundError:
        return {}
    return {r['good_code']: float(r['bag_kg']) for r in rows if r.get('bag_kg')}


def rule_breaks(g, kg):
    out = []
    if len(g['name'] or '') >= NAME_MAX:
        out.append(f'Name {len(g["name"])} chars (limit {NAME_MAX})')
    if g['vat']:
        out.append('VAT on')
    if '303' in g['units'] or g['code'] in kg:
        if g['unit'] != '303':
            out.append(f'kg good: main unit {g["unit"]}, should be 303')
        if not g['is_weight']:
            out.append('kg good: IsWeight off')
        if g['alt_unit'] != '003':
            out.append(f'kg good: alt unit {g["alt_unit"] or "none"}, should be 003')
        want, have = kg.get(g['code']), g['coefs'].get('003')
        if want and have != want:
            out.append(f'kg good: 003 coefficient {have}, bag is {want:g} kg')
    return out


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    out = sys.argv[sys.argv.index('--out') + 1] if '--out' in sys.argv else None
    goods = parse_goods(args[0])
    prices = parse_prices(args[1])
    rows = []
    for g in goods:
        if g['type'] != '1':          # services (delivery …) carry no barcode
            continue
        for u in g['units']:
            bc = g['barcodes'].get(u, [])
            if u == g['unit'] and g['main_barcode'] and g['main_barcode'] not in bc:
                bc = bc + [g['main_barcode']]
            pr = prices.get((g['code'], u))
            if not bc or not pr:
                rows.append([g['code'], u, UNIT_NAMES.get(u, u), len(g['units']),
                             'no' if not bc else ', '.join(bc),
                             'no' if not pr else pr, g['name'], g['long']])
    kg = bag_kg()
    breaks = [[g['code'], b, g['name'], g['long']] for g in goods
              if g['type'] == '1' for b in rule_breaks(g, kg)]
    multi = [g for g in goods if len(g['units']) > 1]
    print(f'goods {len(goods)}, with 2+ units {len(multi)}, price rows {len(prices)}')
    print(f'good+unit rows missing something: {len(rows)}')
    print(f'  missing barcode: {sum(r[4] == "no" for r in rows)}')
    print(f'  missing price:   {sum(r[5] == "no" for r in rows)}')
    print(f'  goods involved:  {len({r[0] for r in rows})}')
    kind = lambda b: re.sub(r'[\d.]+', 'N', b[1])
    print(f'hard-rule breaks: {len(breaks)} on {len({b[0] for b in breaks})} goods')
    for k in sorted({kind(b) for b in breaks}):
        print(f'  {sum(kind(b) == k for b in breaks):4}  {k}')
    if out:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Missing'
        ws.append(['Code', 'Unit', 'Unit name', 'Units on good', 'Barcode', 'Price',
                   'Name (hy)', 'LongName'])
        for r in rows:
            ws.append(r)
        ws = wb.create_sheet('Rules')
        ws.append(['Code', 'Rule broken', 'Name (hy)', 'LongName'])
        for b in breaks:
            ws.append(b)
        wb.save(out)
        print('wrote', out)


if __name__ == '__main__':
    main()
