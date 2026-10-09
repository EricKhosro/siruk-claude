#!/usr/bin/env python3
"""Export the "per kg" (unit 303) prices from `csv/Siruk-hx.xlsx` PRICES as an xlsx.

For every PRICES row whose col C (Չափման միավոր) is `303`:

  Name          <- the <Name> of the goods whose <Code> is the row's col B
  Siruk Name    <- that goods' <LongName>
  PLU           <- the goods <Code> (col B)
  price per kg  <- the row's price (col F)

The goods names come from `csv/all-xml-scale.xml`, the same file the scale and
`build-armsoft-import.py` read.

Usage:
    .venv-img/bin/python scripts/scale-303-prices.py [--out csv/Siruk-hx-303-prices.xlsx]
"""
import argparse, pathlib, re
from collections import OrderedDict
from html import unescape

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
XLSX = ROOT / 'csv/Siruk-hx.xlsx'
XML = ROOT / 'csv/all-xml-scale.xml'


def text(el, tag):
    m = re.search(rf'<{tag}\b[^>]*/>|<{tag}>(.*?)</{tag}>', el, re.S)
    if not m or m.group(0).endswith('/>'):
        return None
    v = unescape((m.group(1) or '').strip())
    return v or None


def goods_names():
    """Code -> (Name, LongName) for every <Goods> block, in XML order."""
    xml = XML.read_text(encoding='utf-8-sig')
    out = OrderedDict()
    for g in re.findall(r'<Goods[ >].*?</Goods>', xml, re.S):
        code = text(g.split('<GoodQuantityUnits', 1)[0], 'Code')
        if code:
            out[code] = (text(g, 'Name'), text(g, 'LongName'))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default='csv/Siruk-hx-303-prices.xlsx')
    args = ap.parse_args()

    names = goods_names()
    ws_in = openpyxl.load_workbook(XLSX, data_only=True)['PRICES']

    rows = []
    for r in ws_in.iter_rows(min_row=3, values_only=True):
        code, unit, price = r[1], r[2], r[5]
        if str(unit).strip() != '303':
            continue
        name, longname = names.get(str(code).strip(), (None, None))
        rows.append([name, longname, code, price])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'PER KG'
    ws.append(['Name', 'Siruk Name', 'PLU', 'price per kg'])
    for row in rows:
        ws.append(row)
    out = ROOT / args.out
    wb.save(out)
    print(f'{len(rows)} rows -> {out}')


if __name__ == '__main__':
    main()
