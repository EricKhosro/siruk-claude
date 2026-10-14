#!/usr/bin/env python3
"""Find live products that are probably variants of ONE product — review only, writes nothing live.

    uv run --with openpyxl python scripts/find-merge-candidates.py --run runs/<date>
      -> <run>/merge-candidates.xlsx (+ merge-candidates.json for plan-variant-merge.py)

Why it exists: CLAUDE.md rule 9 — a flavour, dose band, colour or size inside a product Name is
the symptom of a mis-split product, and since siruk-web 2026-10-09 a mis-split line also floods
its siblings' "Related items" (CLAUDE.md 14). find-duplicate-products.py catches the Trixie
accessory case; food / treats / pharmacy siblings need the name-minus-axis signal below.

How a group forms (all must hold):
  * same brand, same product type, same category set, not discontinued;
  * the same name once every AXIS word is removed — flavour / texture / colour values from
    reference/attribute-values.json, dose bands, letter and measured sizes, pack sizes, "with",
    "and", "in" … (lifestage, breed size, diet and health words are NOT removed, so they keep
    lines apart);
  * no SPLITTING attribute disagrees — the product type's attribute-role values for lifestage,
    breed-size, special-diet, food-form, product-form, health-feature, packaging, material
    (an option-role attribute may differ: that is the variant axis);
  * Trixie's own catalogue (trixie.de page / trixie.shop product, .siruk-cache/dedup maps) does not
    put the members on different products.
Groups that pass the name test but fail a splitting attribute or the Trixie check are listed on a
second sheet ("Kept apart") with the reason — so the decision is visible, not hidden.

Confidence: high = the merged variants are already told apart by an option value or size;
medium = they would collide until an option value (usually flavour) is set — the axis words found
in the names are shown as the proposed value. The workbook has a Decision column for the user.
"""
import json, os, re, sys
from collections import defaultdict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LIVE = os.path.join(ROOT, ".siruk-cache", "live")
SPLITTING = ("lifestage", "breed-size", "special-diet", "food-form", "product-form",
             "health-feature", "packaging", "material")
AXIS_ATTRS = ("flavor", "texture", "color-family")
FILLER = {"with", "and", "&", "in", "+", "of", "a", "the", "flavour", "flavor", "taste", "recipe",
          "pieces", "piece", "sauce", "gravy", "jelly", "pate", "paté",
          "chunks", "fillets", "mousse", "shredded", "minced", "stew", "broth", "vegetables", "veggies",
          "dark", "neon", "colours", "colors", "various", "assorted"}
# Words that can name a LINE (Fresh, Fruit, Rich, Light, Mix …) are deliberately not here.
SIZE_RX = [r"\b\d+([.,]\d+)?\s*[–-]\s*\d+([.,]\d+)?\s*(kg|cm|mm|m|g)\b", r"\b(up to|over)\s+\d+\s*kg\b",
           r"\b\d+([.,]\d+)?\s*(kg|g|ml|l|cm|mm|m|pcs\.?|pc|bags?)\b", r"\b\d+\s*[×x*]\s*\d+([.,]\d+)?\b",
           r"ø\s*\d+", r"\b(xxs|xs|s|m|l|xl|xxl)(\s*[–-]\s*(xxs|xs|s|m|l|xl|xxl))?\b"]


def words_of(label):
    return [w for w in re.split(r"[^a-z0-9äöüéè']+", label.lower()) if w]


def load():
    ids = json.load(open(os.path.join(LIVE, "products", "index.json")))
    prods = {i: json.load(open(os.path.join(LIVE, "products", f"{i}.json"))) for i in ids
             if os.path.exists(os.path.join(LIVE, "products", f"{i}.json"))}
    menu = json.load(open(os.path.join(ROOT, "reference", "attribute-values.json")))
    types = {t["id"]: t for t in json.load(open(os.path.join(ROOT, "reference", "product-types.json")))}
    cats = {}

    def walk(n, trail):
        cats[n["id"]] = " > ".join(trail + [n["name"]])
        for c in n.get("children") or []:
            walk(c, trail + [n["name"]])
    for n in json.load(open(os.path.join(LIVE, "categories.json"))):
        walk(n, [])
    brands = {}
    try:
        if os.environ.get("NO_API") == "1":
            raise RuntimeError("brand names skipped")
        sys.path.insert(0, os.path.join(ROOT, "scripts"))
        from siruk_payload import api
        brands = {b["id"]: b["name"] for b in api("GET", "/brands?forProducts=true").get("data", [])}
    except Exception:
        cached = os.path.join(LIVE, "brand-names.json")     # written from the public API when NO_API=1
        if os.path.exists(cached):
            brands = {int(k): v for k, v in json.load(open(cached)).items()}
    trixie = {}
    for f, key in (("trixie-map.json", "page"), ("shop-map.json", "shop")):
        path = os.path.join(ROOT, ".siruk-cache", "dedup", f)
        if os.path.exists(path):
            m = json.load(open(path))
            trixie[key] = m.get("a2p", m) if isinstance(m, dict) else {}
    return prods, menu, types, cats, brands, trixie


def build(prods, menu, types, cats, brands, trixie):
    attr_id = {code: str(a["attribute_id"]) for code, a in menu.items() if isinstance(a, dict) and "attribute_id" in a}
    label_of = {}
    for code, a in menu.items():
        if isinstance(a, dict):
            for lbl, vid in (a.get("values") or {}).items():
                label_of[(attr_id.get(code), vid)] = lbl
    vocab = set()
    for code in AXIS_ATTRS:
        for lbl in (menu.get(code) or {}).get("values") or {}:
            vocab |= set(words_of(lbl))
    vocab |= {"pate", "paté", "chunks", "gravy", "jelly"}
    axis_words = set(FILLER) | vocab
    for w in ("adult", "puppy", "kitten", "junior", "senior", "mini", "maxi", "medium", "large", "small",
              "giant", "sterilised", "sterilized", "hairball", "indoor", "light", "urinary", "renal",
              "dog", "dogs", "cat", "cats", "grain", "free", "low", "high", "protein", "natural",
              "original", "classic", "mini", "soft", "tender", "steel", "silver", "meat"):
        axis_words.discard(w)     # line words, never axis words
        vocab.discard(w)

    grammar = {"with", "and", "&", "in", "+", "of", "a", "the"}
    vocab_words = vocab - grammar - {"with", "and"}   # real flavour / texture / colour words only

    def stem(name):
        """The line part of a name. First choice: everything before the first axis word ("Gift Soft
        Sticks Adult Cat | Rabbit with Sage") — that also covers garnish words our flavour list lacks
        (sage, rosehips). When the name starts with the axis ("Ostrich Meat Medallions") or the
        prefix is a single word, every axis word is removed instead."""
        n = name.lower()
        for rx in SIZE_RX:
            n = re.sub(rx, " ", n)
        ws = words_of(n)
        prefix = []
        for w in ws:
            if w in vocab_words:          # "with" alone never cuts: "Collar with Rhinestone" is its own line
                break
            prefix.append(w)
        while prefix and prefix[-1] in grammar:
            prefix.pop()
        if len(prefix) >= 2:
            return " ".join(prefix)
        return " ".join(w for w in ws if w not in axis_words)

    def values(p, code):
        aid = attr_id.get(code)
        out = set()
        for v in p["variants"]:
            out |= {label_of.get((aid, i), str(i)) for i in (v.get("attribute_values") or {}).get(aid) or []}
        return out

    def roles(p):
        return {a.get("code"): a["role"] for a in (types.get(p.get("attribute_family_id")) or {}).get("attributes") or []}

    def tkeys(p):
        keys = set()
        for v in p["variants"]:
            s = (v.get("sku") or "").strip()
            for k, m in trixie.items():
                if m.get(s):
                    keys.add((k, str(m[s])))
        return keys

    blocks = defaultdict(list)
    for p in prods.values():
        if p.get("is_discontinued"):
            continue
        blocks[(p.get("brand_id"), p.get("attribute_family_id"), tuple(sorted(p.get("category_ids") or [])),
                stem(p["name"]))].append(p)

    groups, kept = [], []
    for (bid, tid, cids, st), members in blocks.items():
        if len(members) < 2 or not st:
            continue
        r = roles(members[0])
        reasons = []
        for code in SPLITTING:
            if r.get(code) != "attribute":
                continue
            vals = [frozenset(values(p, code)) for p in members]
            set_vals = {v for v in vals if v}
            if len(set_vals) > 1:
                reasons.append(f"{code} differs: " + " / ".join(sorted(", ".join(sorted(v)) for v in set_vals)))
        tk = [tkeys(p) for p in members]
        kinds = defaultdict(set)
        for keys in tk:
            for k, v in keys:
                kinds[k].add(v)
        if any(len(v) > 1 for v in kinds.values()):    # Trixie's own catalogue splits them
            reasons.append("Trixie sells them as separate products (" + ", ".join(
                f"{k}: {len(v)}" for k, v in kinds.items()) + ")")
        option_codes = [c for c, role in r.items() if role == "option"]
        sigs, collide = {}, []
        for p in members:
            for v in p["variants"]:
                sig = (tuple(sorted((c, tuple(sorted(values({"variants": [v]}, c)))) for c in option_codes)),
                       v.get("measure_type"), v.get("content"), v.get("pack_count") or 1)
                if sig in sigs and sigs[sig] != p["id"]:
                    collide.append((sigs[sig], p["id"]))
                sigs.setdefault(sig, p["id"])
        survivor = max(members, key=lambda p: (len(p["variants"]), -p["id"]))
        entry = {"brand": brands.get(bid, str(bid)), "type": (types.get(tid) or {}).get("name", str(tid)),
                 "categories": " | ".join(cats.get(c, str(c)) for c in cids), "stem": st,
                 "survivor": survivor["id"], "members": [], "reasons": reasons,
                 "confidence": "medium" if collide else "high",
                 "collisions": sorted({tuple(sorted(c)) for c in collide})}
        for p in sorted(members, key=lambda p: p["id"]):
            axis = sorted({w for w in words_of(p["name"]) if w in axis_words and w not in FILLER})
            entry["members"].append({
                "id": p["id"], "name": p["name"], "slug": p["slug"],
                "variants": [{"id": v["id"], "sku": v["sku"], "label": v.get("size_label") or v.get("name"),
                              "price": v.get("price"), "stock": v.get("available_quantity")} for v in p["variants"]],
                "options": {c: sorted(values(p, c)) for c in option_codes if values(p, c)},
                "axis_words_in_name": axis,
                "splitting": {c: sorted(values(p, c)) for c in SPLITTING if r.get(c) == "attribute" and values(p, c)},
                "trixie": sorted(f"{k}:{v}" for k, v in tkeys(p))})
        (kept if reasons else groups).append(entry)
    order = lambda g: (g["confidence"] != "high", g["brand"], g["stem"])
    return sorted(groups, key=order), sorted(kept, key=order)


def line_name(g, prods_by_id):
    """A proposed line name: the survivor's own words that make up the group's stem, in its casing
    ("Adult Cat Paté Beef" + stem "adult cat" → "Adult Cat"). The user can overwrite it."""
    keep = set(g["stem"].split())
    out, seen = [], set()
    for tok in re.split(r"\s+", prods_by_id[g["survivor"]]["name"]):
        w = (words_of(tok) or [""])[0]
        if w in keep and w not in seen:
            out.append(tok.strip(","))
            seen.add(w)
    return " ".join(out) or prods_by_id[g["survivor"]]["name"]


def write_xlsx(path, groups, kept, prods):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    pid = {p["id"]: p for p in prods.values()}
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    hi = sum(1 for g in groups if g["confidence"] == "high")
    lines = [
        ("Merge review — products that are probably one product", True),
        ("Nothing has been merged. Fill the Decision column on 'Merge candidates' (Merge / Keep apart / Unsure) and hand the file back.", False),
        ("", False),
        (f"Merge candidates: {len(groups)} groups, {sum(len(g['members']) for g in groups)} products "
         f"→ {sum(len(g['members']) - 1 for g in groups)} fewer products if all merged. High confidence: {hi}.", False),
        (f"Kept apart (for information): {len(kept)} look-alike groups the rules keep separate, with the reason.", False),
        ("", False),
        ("How a group forms: same brand, product type and categories, and the same name once the flavour / texture / colour / size words are removed.", False),
        ("Never grouped: lines that differ by lifestage, breed size, diet, health function, packaging or material (those are separate products by our rules).", False),
        ("Confidence high: the variants are already told apart by an option value or pack size, so a merge needs no new data.", False),
        ("Confidence medium: some variants would look identical in the selector until an option value (usually Flavour) is set — 'Axis words in name' is the proposed value.", False),
        ("Survivor = the product that keeps its id, page and reviews; the others are folded into it as variants.", False),
        ("Links open the live storefront page of each variant.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        ws.cell(i, 1, t).font = Font(bold=b, size=13 if b else 11)
    ws.column_dimensions["A"].width = 150

    cols = ["Group", "Confidence", "Decision", "Proposed line name", "Survivor", "Product id", "Current name",
            "Words only this name has", "Brand", "Product type", "Categories",
            "Variants (label · sku · price · stock)", "Option values set", "Axis words in name",
            "Would collide with", "Storefront link"]
    ws = wb.create_sheet("Merge candidates")
    ws.append(cols)
    dv = DataValidation(type="list", formula1='"Merge,Keep apart,Unsure"', allow_blank=True)
    ws.add_data_validation(dv)
    fills = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="EEF3FB")]
    for gi, g in enumerate(groups, 1):
        proposed = line_name(g, pid)
        common = set.intersection(*[set(words_of(m["name"])) for m in g["members"]])
        for m in g["members"]:
            coll = [b if a == m["id"] else a for a, b in g["collisions"] if m["id"] in (a, b)]
            v0 = m["variants"][0]
            ws.append([gi, g["confidence"], "", proposed, "yes" if m["id"] == g["survivor"] else "",
                       m["id"], m["name"], ", ".join(w for w in dict.fromkeys(words_of(m["name"])) if w not in common),
                       g["brand"], g["type"], g["categories"],
                       "\n".join(f"{v['label']} · {v['sku']} · {v['price']} · {v['stock']}" for v in m["variants"]),
                       "; ".join(f"{c}: {', '.join(v)}" for c, v in m["options"].items()),
                       ", ".join(m["axis_words_in_name"]),
                       ", ".join(str(c) for c in coll),
                       f"https://siruk.am/product/{m['slug']}/dp/{v0['id']}/"])
            row = ws.max_row
            dv.add(ws.cell(row, 3))
            for c in range(1, len(cols) + 1):
                ws.cell(row, c).fill = fills[gi % 2]
                ws.cell(row, c).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row, 16).hyperlink = ws.cell(row, 16).value
    widths = [7, 11, 12, 34, 9, 10, 40, 22, 14, 13, 34, 40, 30, 18, 12, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Kept apart")
    ws.append(["Group", "Why kept apart", "Decision", "Product id", "Current name", "Brand", "Product type",
               "Categories", "Differing values", "Storefront link"])
    dv2 = DataValidation(type="list", formula1='"Keep apart,Merge,Unsure"', allow_blank=True)
    ws.add_data_validation(dv2)
    for gi, g in enumerate(kept, 1):
        for m in g["members"]:
            ws.append([gi, "; ".join(g["reasons"]), "", m["id"], m["name"], g["brand"], g["type"], g["categories"],
                       "; ".join(f"{c}: {', '.join(v)}" for c, v in m["splitting"].items()) +
                       (f"; trixie {', '.join(m['trixie'])}" if m["trixie"] else ""),
                       f"https://siruk.am/product/{m['slug']}/dp/{m['variants'][0]['id']}/"])
            dv2.add(ws.cell(ws.max_row, 3))
            for c in range(1, 11):
                ws.cell(ws.max_row, c).fill = fills[gi % 2]
                ws.cell(ws.max_row, c).alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate([7, 40, 12, 10, 42, 14, 13, 34, 40, 40], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "D2"
    wb.save(path)


def main(argv):
    if "--run" not in argv:
        sys.exit(__doc__)
    run = os.path.join(ROOT, argv[argv.index("--run") + 1])
    os.makedirs(run, exist_ok=True)
    prods, menu, types, cats, brands, trixie = load()
    groups, kept = build(prods, menu, types, cats, brands, trixie)
    json.dump({"groups": groups, "kept_apart": kept}, open(os.path.join(run, "merge-candidates.json"), "w"),
              ensure_ascii=False, indent=1)
    write_xlsx(os.path.join(run, "merge-candidates.xlsx"), groups, kept, prods)
    print(f"{len(groups)} merge groups ({sum(1 for g in groups if g['confidence'] == 'high')} high), "
          f"{sum(len(g['members']) for g in groups)} products; {len(kept)} kept-apart groups")


if __name__ == "__main__":
    main(sys.argv[1:])
