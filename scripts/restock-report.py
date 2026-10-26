#!/usr/bin/env python3
"""Stock left vs sold per ArmSoft good, least stock first — the restock sheet.

    uv run --with openpyxl python scripts/restock-report.py <turnover.xlsx> \
        [--prices <armsoft-prices.xlsx>] [--stock <armsoft-stock.xlsx>] --out restock.xlsx

<turnover.xlsx>  ArmSoft "Ապրանքների շրջանառություն" (navigator `շրջանառ` →
                 Պահեստ), period from 01/01/01, "Բացված ըստ գործողությունների"
                 ticked, Արտահանել Excel. One column per operation type.
--prices         ArmSoft price list (fallback sale price for goods not on Siruk).
--stock          ArmSoft stock report (fallback cost, ex VAT → +20 %).

Sold = Վաճառքներ + Վաճառք (Մանրածախ) − Վերադարձներ. Warehouse issue orders
(Պահեստի ելքի օրդեր) are NOT sales: they were stock corrections (user,
2026-10-26). Buying / selling price and the link come from the live Siruk
catalogue (.siruk-cache/live/products, refresh with scripts/dump-products.py),
matched on the variant's accounting_code. Per-kg goods are counted in kg, so
their prices are per kg too (the 1 kg / 100 g variant, else bag ÷ bag kg).
Duplicate goods (state/armsoft-duplicate-codes.csv: an old copy of a kept
good, same EAN) are folded into the kept code's row — their sales and stock
add up there, and the old code gets no row of its own.
"""
import argparse, csv, glob, json, os, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
UNIT_CODE = {"կգ": "303", "հատ": "001", "փաթեթ": "004", "պարկ": "003", "տուփ": "002"}
# goods kept in ArmSoft on purpose but not sold on Siruk under their own code
KNOWN_NOTE = {**{c: "RC 12×85 g box good: its stock moved to the pouch goods (knowledge/goods-setups.md)"
                 for c in ("1466", "1467", "1468", "1469", "1470", "1471", "1472")},
              "1442": "unidentified good: brand + EAN asked of the seller (state/open-items.csv)"}


def load_duplicates():
    path = os.path.join(ROOT, "state", "armsoft-duplicate-codes.csv")
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as f:
        return {r["old_code"]: r["keep_code"] for r in csv.DictReader(f)}


def fold_duplicates(goods, dup):
    """Add each old duplicate good's sold / received / stock to its kept good."""
    by_code = {g["code"]: g for g in goods}
    out = []
    for g in goods:
        keep = by_code.get(dup.get(g["code"]))
        if not keep:
            out.append(g)
            continue
        for k in ("sold", "received", "stock"):
            keep[k] = round(keep[k] + g[k], 3)
        if g["sold"] or g["stock"]:
            keep.setdefault("merged", []).append(f"{g['code']} (sold {g['sold']:g}, stock {g['stock']:g})")
        else:
            keep.setdefault("merged", []).append(g["code"])
    return out


def col(header, name):
    return header.index(name)


def load_turnover(path):
    rows = list(openpyxl.load_workbook(path, read_only=True).active.iter_rows(values_only=True))
    h = list(rows[2])
    i_ret = col(h, "Վերադարձներ/Քանակ մուտք")
    i_s1, i_s2 = col(h, "Վաճառքներ/Քանակ ելք"), col(h, "Վաճառք (Մանրածախ)/Քանակ ելք")
    i_end = col(h, "Վերջնական քանակ")
    i_in = [i for i, x in enumerate(h) if x and x.endswith("Քանակ մուտք")]
    out = []
    for r in rows[3:]:
        if not r or not r[0] or r[0] == "Ընդամենը":
            continue
        n = lambda i: float(r[i] or 0)
        out.append({"code": str(r[0]), "name": r[1], "unit": r[2] or "",
                    "sold": round(n(i_s1) + n(i_s2) - n(i_ret), 3),
                    "received": round(sum(n(i) for i in i_in) + float(r[3] or 0), 3),
                    "stock": round(n(i_end), 3)})
    return out


def load_prices(path):
    """(code, unit) -> latest sale price (type 01)."""
    best = {}
    if not path:
        return best
    for r in list(openpyxl.load_workbook(path, read_only=True).active.iter_rows(values_only=True))[2:]:
        if not r or len(r) < 11 or r[5] != "01":
            continue
        k = (str(r[1]), str(r[7]))
        if k not in best or r[0] >= best[k][0]:
            best[k] = (r[0], r[10])
    return {k: v[1] for k, v in best.items()}


def load_stock_cost(path):
    cost = {}
    if not path:
        return cost
    for r in list(openpyxl.load_workbook(path, read_only=True).active.iter_rows(values_only=True))[2:]:
        if r and r[0] and r[5]:
            cost[str(r[0])] = float(r[5])
    return cost


def load_siruk():
    brands = {}
    try:
        brands = {b["id"]: b["name"] for b in json.load(open(os.path.join(ROOT, ".siruk-cache", "brands.json"), encoding="utf-8"))}
    except Exception:
        pass
    by_code = {}
    for f in glob.glob(os.path.join(ROOT, ".siruk-cache", "live", "products", "*.json")):
        if f.endswith("index.json"):
            continue
        p = json.load(open(f, encoding="utf-8"))
        p = p.get("product", p)
        for v in p.get("variants", []):
            c = (v.get("accounting_code") or "").strip()
            if c:
                by_code.setdefault(c, []).append((p, v))
    return by_code, brands


def siruk_row(matches, per_kg, brands, bag_kg_from_name=None):
    """Pick the main variant and compute the unit cost / price in ArmSoft's stock unit."""
    def kind(v):
        s = v.get("sku") or ""
        return "kg" if s.endswith("-KG") else "100g" if s.endswith("-100G") else "single" if s.endswith("-1") else "main"
    main = [m for m in matches if kind(m[1]) == "main"] or matches
    p, v = main[0]
    brand = brands.get(p.get("brand_id"), "")
    label = v.get("name") or v.get("size_label")   # the label carries the flavour; size_label is size only
    name = f"{brand} {p['name']}".strip() + (f", {label}" if label else "")
    link = f"https://siruk.am/product/{p['slug']}/dp/{v['id']}/"
    cost, price = v.get("cost_price"), v.get("price")
    if per_kg:
        kg = (v.get("content") or 0) / 1000 * (v.get("pack_count") or 1) if v.get("measure_type") == "mass" else 0
        loose = {kind(m[1]): m[1] for m in matches}
        if "kg" in loose:
            price, cost = loose["kg"].get("price"), loose["kg"].get("cost_price") or (cost / kg if kg and cost else None)
        elif "100g" in loose:
            price = loose["100g"]["price"] * 10 if loose["100g"].get("price") else None
            cost = cost / kg if kg and cost else (loose["100g"].get("cost_price") or 0) * 10 or None
        else:
            kg = kg or bag_kg_from_name   # litter: a volume on Siruk, kg in ArmSoft
            if kg:
                price, cost = (price / kg if price else None), (cost / kg if cost else None)
            else:
                price = cost = None
        bag_kg = kg or None
    else:
        bag_kg = None
    return {"name": name, "link": link, "cost": cost, "price": price, "bag_kg": bag_kg,
            "note": "" if len(main) == 1 else f"{len(main)} Siruk variants share this code"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("turnover")
    ap.add_argument("--prices")
    ap.add_argument("--stock")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    goods = fold_duplicates(load_turnover(a.turnover), load_duplicates())
    prices, acost = load_prices(a.prices), load_stock_cost(a.stock)
    siruk, brands = load_siruk()

    rows = []
    for g in goods:
        if g["received"] <= 0 and g["sold"] == 0 and g["stock"] == 0:
            continue                       # never stocked
        per_kg = g["unit"] == "կգ"
        m = re.search(r"(\d+(?:[.,]\d+)?)\s*կգ", g["name"] or "")
        s = siruk_row(siruk[g["code"]], per_kg, brands,
                      float(m.group(1).replace(",", ".")) if m else None) if g["code"] in siruk else None
        notes = []
        cost = s["cost"] if s else None
        price = s["price"] if s else None
        if not s:
            notes.append(KNOWN_NOTE.get(g["code"], "not on Siruk"))
        elif s["note"]:
            notes.append(s["note"])
        if g.get("merged"):
            notes.append("incl. old duplicate code " + ", ".join(g["merged"]))
        if price is None:
            price = prices.get((g["code"], UNIT_CODE.get(g["unit"], "001")))
            if price is not None:
                notes.append("sale price from ArmSoft")
        if cost is None and g["code"] in acost:
            cost = round(acost[g["code"]] * 1.2, 2)
            notes.append("cost = ArmSoft cost + 20% VAT")
        margin = round(price - cost, 2) if price and cost else None
        pct = margin / price if margin is not None and price else None
        if g["stock"] < 0:
            notes.append("negative stock: sold more than received")
        rows.append([g["code"], s["name"] if s else "", g["name"], g["unit"], g["stock"],
                     round(g["stock"] / s["bag_kg"], 2) if s and s["bag_kg"] and per_kg else None,
                     g["sold"], cost, price, margin, pct,
                     s["link"] if s else "", "; ".join(notes)])

    rows.sort(key=lambda r: (r[4], -r[6], r[0]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Restock"
    head = ["ArmSoft code", "Product (Siruk)", "ArmSoft name", "Unit", "Stock left", "≈ bags left",
            "Sold", "Buying price", "Selling price", "Margin (AMD)", "Margin %", "Siruk link", "Note"]
    ws.append(head)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    red, amber = PatternFill("solid", fgColor="F8CBAD"), PatternFill("solid", fgColor="FFE699")
    for i, r in enumerate(rows, start=2):
        fill = red if r[4] <= 0 else amber if (r[6] > 0 and r[4] <= r[6]) else None
        if fill:
            for c in ws[i][:7]:
                c.fill = fill
        for j in (8, 9, 10):
            ws.cell(i, j).number_format = "#,##0"
        ws.cell(i, 11).number_format = "0.0%"
        for j in (5, 6, 7):
            ws.cell(i, j).number_format = "#,##0.###"
        if r[11]:
            ws.cell(i, 12).hyperlink = r[11]
            ws.cell(i, 12).font = Font(color="0563C1", underline="single")
    for j, w in enumerate([10, 46, 40, 7, 10, 10, 9, 12, 12, 12, 9, 50, 34], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    n = ws.max_row + 2
    ws.cell(n, 1, "Sold = till sales (Վաճառքներ + Վաճառք մանրածախ) minus returns; warehouse issue orders are stock corrections and not counted.").font = Font(italic=True)
    ws.cell(n + 1, 1, "Per-kg goods (unit կգ): stock, sold and prices are per kg. Red = out of stock; yellow = stock left ≤ sold so far.").font = Font(italic=True)
    wb.save(a.out)
    sold = sum(1 for r in rows if r[6] > 0)
    print(f"{len(rows)} goods ({sold} sold, {sum(1 for r in rows if r[4] <= 0)} at or below 0) -> {a.out}")
    print("not on Siruk:", sum(1 for r in rows if not r[11]), "| no price:", sum(1 for r in rows if r[8] is None), "| no cost:", sum(1 for r in rows if r[7] is None))


if __name__ == "__main__":
    main()
