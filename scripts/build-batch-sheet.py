#!/usr/bin/env python3
"""Step 0 of an import: the batch's invoices -> one Excel, the batch's source of truth.

    uv run --with openpyxl python3 scripts/build-batch-sheet.py batches/<batch> [--force]
    uv run --with openpyxl python3 scripts/build-batch-sheet.py batches/<batch> --from-xlsx

Reads batches/<batch>/lines/*.json, one per invoice, transcribed by hand from
batches/<batch>/invoices/ (any language, any headers; the shape is in
reference/invoice-intake.md). Then:

  - checks every line (qty x unit price = line total) and every invoice
    (sum of lines = printed total, VAT basis respected)
  - cost per retail unit = unit price incl. VAT (user, 2026-10-14), / units per
    item for a case sold singly, x the invoice's fx_rate when not AMD (the user
    gives the rate; never guessed: no rate -> cost blank + a question)
  - folds lines into one row per article code (a codeless line keys by name),
    with qty summed and a cost that differs between invoices flagged
  - guesses Brand / Species / Category with the batch-1 rules
    (build-product-catalogue.py); Brand is a hint only (the EAN/hafo decides)

Writes batches/<batch>/batch.xlsx (sheets Products, Lines, Invoices, Questions)
and batches/<batch>/products.csv (Shape C, what prepare-run.py / /add-products read).
It refuses to overwrite an existing batch.xlsx without --force, because the user may
have edited it. --from-xlsx goes the other way: the (edited) Products sheet ->
products.csv, and prints the open questions.

Exit 0 = nothing flagged, 2 = the Questions sheet has open rows, 1 = input error.
"""
import argparse, collections, csv, glob, importlib.util, json, os, re, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("needs openpyxl: uv run --with openpyxl python3 scripts/build-batch-sheet.py …")

_spec = importlib.util.spec_from_file_location(
    "catalogue", os.path.join(ROOT, "scripts", "build-product-catalogue.py"))
catalogue = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(catalogue)

PRODUCT_COLS = ["Article Code", "Brand", "Product Name (as printed)", "Species", "Category",
                "Buy Price (AMD)", "Qty Received", "Sale Price (AMD)", "EAN", "Invoice Date",
                "Source", "Brand Source", "Invoice Lines", "Notes"]
LINE_COLS = ["Invoice", "Page", "Line", "Article Code", "Name (as printed)", "Unit", "Qty",
             "Unit Price (printed)", "Line Total (printed)", "Line VAT", "Units per Item",
             "Currency", "Cost per Unit incl. VAT", "Buy Price (AMD)", "EAN", "Supplier RRP", "Note"]
INVOICE_COLS = ["Invoice", "Supplier", "Invoice No", "Date", "Language", "Currency", "FX Rate",
                "Prices incl. VAT", "VAT %", "Lines", "Sum of Lines", "Printed Subtotal",
                "Printed Total", "Status", "Headers Read As"]
QUESTION_COLS = ["#", "Where", "Kind", "Question", "Answer"]

TOL_ABS, TOL_REL = 1.0, 0.005    # a line or a total may be off by 1 unit of currency or 0.5 %
CASE_RE = re.compile(r"(?<![\d.,])(\d{1,3})\s*[xX×хХ*]\s*\d")   # "12X70G", "6x1,5kg"


def close(a, b, n=1):
    return abs(a - b) <= max(TOL_ABS * n, TOL_REL * abs(b))


def fnum(x):
    if x is None or x == "":
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).replace(" ", "").replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", s):
        s = s.replace(",", "")
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def tidy(x):
    if x is None:
        return None
    x = round(x, 2)
    return int(x) if x == int(x) else x


class Questions:
    def __init__(self):
        self.rows = []

    def add(self, where, kind, text):
        self.rows.append([len(self.rows) + 1, where, kind, text, ""])


def load_invoices(batch):
    paths = sorted(glob.glob(os.path.join(batch, "lines", "*.json")))
    if not paths:
        sys.exit(f"no transcriptions in {batch}/lines/*.json; transcribe the invoices first "
                 "(reference/invoice-intake.md)")
    invs = []
    for p in paths:
        try:
            inv = json.load(open(p))
        except json.JSONDecodeError as e:
            sys.exit(f"{p}: not valid JSON ({e})")
        inv["_json"] = os.path.relpath(p, batch)
        inv.setdefault("file", inv["_json"])
        if not isinstance(inv.get("lines"), list) or not inv["lines"]:
            sys.exit(f"{p}: no lines")
        invs.append(inv)
    return invs


def check_invoice(inv, q):
    """Per-line costs + the invoice reconcile. Returns (lines out, invoice row)."""
    f = inv["file"]
    cur = (inv.get("currency") or "AMD").upper()
    fx = fnum(inv.get("fx_rate"))
    incl = inv.get("prices_include_vat")
    rate = fnum(inv.get("vat_rate")) or 0.0
    if incl is None:
        q.add(f, "vat", "The invoice does not say whether its prices include VAT; set "
                        "prices_include_vat in the transcription.")
        incl = True
    if cur != "AMD" and not fx:
        q.add(f, "currency", f"Invoice is in {cur}: what AMD rate (or landed cost) should we use? "
                             "Costs on this invoice are blank until set.")
    out, total_net, total_vat = [], 0.0, 0.0
    for i, ln in enumerate(inv["lines"], 1):
        where = f"{f} " + (f"p{ln['page']} " if ln.get("page") else "") + f"l{ln.get('line', i)}"
        qty, price, ltot = fnum(ln.get("qty")), fnum(ln.get("unit_price")), fnum(ln.get("line_total"))
        lvat = fnum(ln.get("line_vat"))
        upi = fnum(ln.get("units_per_item")) or 1.0
        name = (ln.get("name") or "").strip()
        code = (ln.get("code") or "").strip()
        if not code:
            q.add(where, "identity", f"No article code on the line \"{name}\"; identified by name "
                                     "later (CLAUDE.md rule 6).")
        if ln.get("unreadable"):
            q.add(where, "unreadable", f"Could not read {', '.join(ln['unreadable'])} on \"{name}\"; "
                                       "please check the original.")
        if ln.get("note"):
            q.add(where, "note", ln["note"])
        if qty is None or price is None:
            q.add(where, "missing", f"Qty or unit price missing on \"{name}\".")
        elif ltot is not None and not close(qty * price, ltot):
            q.add(where, "line-total", f"{qty:g} × {price:g} = {qty * price:,.2f} but the line total "
                                       f"reads {ltot:,.2f}; a digit is misread.")
        if ltot is None and qty is not None and price is not None:
            ltot = qty * price
        if upi == 1 and CASE_RE.search(name) and not ln.get("units_per_item"):
            q.add(where, "case", f"\"{name}\" looks like a case. Is {price:g} the price of the whole case "
                                 "(set units_per_item) or of one unit?")
        unit_cost = None
        if price is not None:
            if incl:
                unit_cost = price
            elif lvat is not None and qty:
                unit_cost = (ltot + lvat) / qty
            else:
                unit_cost = price * (1 + rate / 100)
            unit_cost /= upi
        amd = unit_cost if cur == "AMD" else (unit_cost * fx if unit_cost is not None and fx else None)
        total_net += ltot or 0
        total_vat += lvat or 0
        out.append({"invoice": f, "page": ln.get("page"), "line": ln.get("line", i), "code": code,
                    "name": name, "unit": ln.get("unit", ""), "qty": qty, "price": price,
                    "ltot": ltot, "lvat": lvat, "upi": upi, "cur": cur, "unit_cost": unit_cost,
                    "amd": amd, "ean": str(ln.get("barcode") or "").strip(),
                    "rrp": fnum(ln.get("supplier_rrp")), "date": inv.get("date", ""), "where": where})

    printed, sub = fnum(inv.get("printed_total")), fnum(inv.get("printed_subtotal"))
    n = len(out)
    if printed is None and sub is None:
        status = "NO PRINTED TOTAL"
        q.add(f, "total", "No printed total was read, so the lines could not be checked against it.")
    else:
        checks = []
        if incl:
            checks.append(("total", total_net, printed if printed is not None else sub))
        else:
            if sub is not None:
                checks.append(("subtotal", total_net, sub))
            if printed is not None:
                gross = total_net + total_vat if total_vat else total_net * (1 + rate / 100)
                checks.append(("total", gross, printed))
        bad = [(k, s, p) for k, s, p in checks if not close(s, p, n)]
        status = "OK" if not bad else "; ".join(f"{k} {s:,.2f} vs printed {p:,.2f} ({s - p:+,.2f})"
                                                 for k, s, p in bad)
        for k, s, p in bad:
            q.add(f, "total", f"Lines add up to {s:,.2f} but the printed {k} is {p:,.2f} "
                              f"({s - p:+,.2f}): a line is misread or missing.")
    hdr = "; ".join(f"{k} → {v}" for k, v in (inv.get("headers") or {}).items())
    row = [f, inv.get("supplier", ""), inv.get("invoice_no", ""), inv.get("date", ""),
           inv.get("language", ""), cur, fx, "yes" if incl else "no", rate, n, tidy(total_net),
           sub, printed, status, hdr]
    return out, row


def fold(lines, q):
    """One product row per article code (codeless lines key by name)."""
    groups = collections.OrderedDict()
    for ln in lines:
        key = ln["code"] or "name:" + re.sub(r"\s+", " ", ln["name"].lower())
        groups.setdefault(key, []).append(ln)
    rows = []
    for key, g in groups.items():
        g.sort(key=lambda l: str(l["date"]))
        last = g[-1]
        name = max((l["name"] for l in g), key=len)
        costs = {tidy(l["amd"]) for l in g if l["amd"] is not None}
        if len(costs) > 1:
            q.add(last["code"] or name, "cost",
                  f"Cost differs between invoices ({', '.join(str(c) for c in sorted(costs))}); "
                  f"using the latest ({tidy(last['amd'])}, {last['invoice']}).")
        eans = {l["ean"] for l in g if l["ean"]}
        if len(eans) > 1:
            q.add(last["code"] or name, "ean", f"More than one barcode for one code: {', '.join(sorted(eans))}.")
        qty = sum((l["qty"] or 0) * l["upi"] for l in g)
        brand, how = catalogue.brand_of(name, last["code"])
        rows.append({
            "Article Code": last["code"], "Brand": brand, "Product Name (as printed)": name,
            "Species": catalogue.species_of(name), "Category": catalogue.category_of(name),
            "Buy Price (AMD)": tidy(last["amd"]), "Qty Received": tidy(qty), "Sale Price (AMD)": "",
            "EAN": next(iter(sorted(eans)), ""), "Invoice Date": g[0]["date"],
            "Source": ", ".join(sorted({l["invoice"] for l in g})), "Brand Source": how,
            "Invoice Lines": "; ".join(l["where"] for l in g), "Notes": ""})
    return rows


def write_sheet(ws, cols, rows, money=()):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF4")
    for r in rows:
        ws.append([r.get(c, "") for c in cols] if isinstance(r, dict) else r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(cols, 1):
        width = max([len(str(c))] + [len(str(ws.cell(row=j, column=i).value or ""))
                                     for j in range(2, min(ws.max_row, 300) + 1)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
        if c in money:
            for j in range(2, ws.max_row + 1):
                ws.cell(row=j, column=i).number_format = "#,##0.##"


def write_csv(path, rows):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=PRODUCT_COLS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def build(batch, force):
    xlsx = os.path.join(batch, "batch.xlsx")
    if os.path.exists(xlsx) and not force:
        sys.exit(f"{xlsx} exists (it may hold your edits). Use --from-xlsx to export it, "
                 "or --force to rebuild it from the transcriptions.")
    q = Questions()
    lines, inv_rows = [], []
    for inv in load_invoices(batch):
        out, row = check_invoice(inv, q)
        lines += out
        inv_rows.append(row)
    products = fold(lines, q)
    for p in products:
        if not p["Brand"]:
            p["Notes"] = "brand not recognised from the name"
        if p["Buy Price (AMD)"] in (None, ""):
            p["Notes"] = (p["Notes"] + "; " if p["Notes"] else "") + "no cost yet (see Questions)"

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Products"
    write_sheet(ws, PRODUCT_COLS, products, money=("Buy Price (AMD)",))
    line_rows = [[l["invoice"], l["page"], l["line"], l["code"], l["name"], l["unit"], l["qty"],
                  l["price"], l["ltot"], l["lvat"], l["upi"], l["cur"], tidy(l["unit_cost"]),
                  tidy(l["amd"]), l["ean"], l["rrp"], ""] for l in lines]
    write_sheet(wb.create_sheet("Lines"), LINE_COLS, line_rows,
                money=("Unit Price (printed)", "Line Total (printed)", "Cost per Unit incl. VAT",
                       "Buy Price (AMD)"))
    write_sheet(wb.create_sheet("Invoices"), INVOICE_COLS, inv_rows,
                money=("Sum of Lines", "Printed Subtotal", "Printed Total"))
    write_sheet(wb.create_sheet("Questions"), QUESTION_COLS, q.rows)
    wb.save(xlsx)
    write_csv(os.path.join(batch, "products.csv"), products)

    ok = sum(1 for r in inv_rows if r[13] == "OK")
    print(f"{len(inv_rows)} invoices ({ok} reconcile), {len(lines)} lines -> {len(products)} products")
    print(f"  {xlsx}\n  {os.path.join(batch, 'products.csv')}")
    print(f"  {len(q.rows)} open questions" + (": " + ", ".join(
        f"{k} {n}" for k, n in collections.Counter(r[2] for r in q.rows).most_common()) if q.rows else ""))
    return 2 if q.rows else 0


def export(batch):
    xlsx = os.path.join(batch, "batch.xlsx")
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["Products"]
    hdr = [c.value for c in ws[1]]
    missing = [c for c in ("Article Code", "Product Name (as printed)", "Buy Price (AMD)") if c not in hdr]
    if missing:
        sys.exit(f"{xlsx} Products sheet lacks {', '.join(missing)}")
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = {h: ("" if v is None else v) for h, v in zip(hdr, r) if h}
        if any(str(v).strip() for v in d.values()):
            rows.append(d)
    write_csv(os.path.join(batch, "products.csv"), rows)
    open_q = [r for r in wb["Questions"].iter_rows(min_row=2, values_only=True)
              if r and r[0] is not None and not (r[4] or "").strip()] if "Questions" in wb.sheetnames else []
    no_cost = [d.get("Article Code") or d.get("Product Name (as printed)") for d in rows
               if fnum(d.get("Buy Price (AMD)")) is None]
    print(f"{len(rows)} products -> {os.path.join(batch, 'products.csv')}")
    if no_cost:
        print(f"  {len(no_cost)} without a cost: {', '.join(map(str, no_cost[:10]))}")
    print(f"  {len(open_q)} questions without an answer")
    return 2 if open_q or no_cost else 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("batch", help="batches/<batch> folder")
    ap.add_argument("--from-xlsx", action="store_true", help="export the edited Products sheet to products.csv")
    ap.add_argument("--force", action="store_true", help="overwrite an existing batch.xlsx")
    a = ap.parse_args()
    if not os.path.isdir(a.batch):
        sys.exit(f"no folder {a.batch}")
    sys.exit(export(a.batch) if a.from_xlsx else build(a.batch, a.force))


if __name__ == "__main__":
    main()
