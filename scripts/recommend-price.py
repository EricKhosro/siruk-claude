#!/usr/bin/env python3
"""Recommend a sale price for an item no rival sells, from comparable products we already price. Never writes.

    uv run --with openpyxl python scripts/recommend-price.py items.json --out final/price-recommendations-<date>.xlsx

CLAUDE.md rule 2c (user, 2026-10-12): the sale price is the highest of hafo / zoovet / nemo for the exact
item. When none of them sells it (and the sibling rule 2a doesn't apply), we do NOT invent a price. We
RECOMMEND one for the user to approve, from comparables:
  * same brand, same animal (category root), same product type; closer is better: same leaf
    category, then same measure type and nearest pack size;
  * only comparables whose own price is a rival price (final/siruk-source-of-truth.csv "Source of price"
    names hafo / zoovet / nemo), never another recommendation;
  * method A, markup: our cost × the median price/cost ratio of the comparables;
  * method B, unit price: the median price per g / ml / piece of same-measure comparables × our pack
    size (shown whenever a same-measure comparable exists; a >25% gap from A is flagged);
  * recommended = method A rounded to 10 AMD (method B shown beside it as a cross-check); must be above
    cost, else flagged.
items.json: [{"key": "45689Tx", "name": "Plush Monster cat toy 12 cm, pink", "brand_id": 8, "animal": "Cat",
              "type": "Toys", "leaf_id": 23, "measure_type": null, "content": null, "cost": 1000}, …]
Inputs: .siruk-cache/live/products/ (scripts/dump-products.py), the live category tree
(.siruk-cache/live/categories-after.json or categories.json), final/siruk-source-of-truth.csv.
Sheets: Recommendations (one row per item, Decision column) · Comparables (the evidence per item).
"""
import csv, json, os, re, statistics, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LIVE = os.path.join(ROOT, ".siruk-cache", "live")
RIVALS = ("hafo.am", "zoovet.am", "nemo.am")


def r10(x):
    return int(round(x / 10.0)) * 10


def tree():
    root, path = {}, {}
    p = os.path.join(LIVE, "categories-after.json")
    data = json.load(open(p if os.path.exists(p) else os.path.join(LIVE, "categories.json")))

    def walk(n, r, trail):
        root[n["id"]] = r
        path[n["id"]] = " > ".join(trail + [n["name"]])
        for c in n.get("children") or []:
            walk(c, r, trail + [n["name"]])
    for n in data.get("data", data) if isinstance(data, dict) else data:
        walk(n, n["name"], [])
    return root, path


def truth():
    out = {}
    for t in csv.DictReader(open(os.path.join(ROOT, "final", "siruk-source-of-truth.csv"), encoding="utf-8-sig")):
        m = re.search(r"/dp/(\d+)", t["Siruk link"])
        if m:
            out[int(m.group(1))] = t
    return out


def comparables(item, products, root, path, tr):
    rows = []
    for p in products:
        if p.get("is_discontinued") or p.get("brand_id") != item["brand_id"]:
            continue
        roots = {root.get(c) for c in p.get("category_ids") or []}
        if item["animal"] not in roots or p.get("attribute_family_name") != item["type"]:
            continue
        for v in p["variants"]:
            t = tr.get(v["id"])
            if not t or not any(r in t["Source of price"] for r in RIVALS):
                continue
            price, cost = float(t["Price (AMD)"] or 0), float(t["Buying price (AMD)"] or 0)
            if price <= 0 or cost <= 0:
                continue
            same_leaf = item.get("leaf_id") in (p.get("category_ids") or [])
            same_measure = bool(item.get("measure_type")) and v.get("measure_type") == item.get("measure_type") \
                and (v.get("content") or 0) > 0
            unit = price / (v["content"] * (v.get("pack_count") or 1)) if same_measure else None
            dist = abs((v.get("content") or 0) - (item.get("content") or 0)) if same_measure and item.get("content") else 9e9
            rows.append({"product_id": p["id"], "product": p["name"], "variant": v.get("size_label") or v.get("name"),
                         "sku": v["sku"], "price": price, "cost": cost, "markup": price / cost, "unit": unit,
                         "same_leaf": same_leaf, "same_measure": same_measure, "dist": dist,
                         "category": " | ".join(path.get(c, "") for c in p.get("category_ids") or []),
                         "source": t["Source of price"].split(" (")[0],
                         "link": f"https://siruk.am/product/{p['slug']}/dp/{v['id']}/"})
    rows.sort(key=lambda r: (not r["same_leaf"], not r["same_measure"], r["dist"]))
    return rows


def recommend(item, comps):
    if not comps:
        return {"recommended": None, "method_a": None, "method_b": None, "n": 0, "note": "no comparable with a rival price"}
    pool = [c for c in comps if c["same_leaf"]] or comps
    mk = statistics.median(c["markup"] for c in pool)
    a = r10(item["cost"] * mk)
    units = [c["unit"] for c in pool if c["unit"]] or [c["unit"] for c in comps if c["unit"]]
    b = r10(statistics.median(units) * item["content"]) if units and item.get("content") else None
    note = f"median markup ×{mk:.2f} over {len(pool)} comparables" + (" in the same category" if pool is not comps else "")
    if len(pool) < 3:
        note += "; LOW CONFIDENCE (fewer than 3 comparables)"
    if b is not None and abs(b - a) / a > 0.25:
        note += f"; unit-price method says {b} — sizes differ, check"
    if a <= item["cost"]:
        note += "; BELOW/AT COST — check"
    return {"recommended": a, "method_a": a, "method_b": b, "n": len(pool), "note": note}


def main(argv):
    items = json.load(open(argv[0]))
    out = argv[argv.index("--out") + 1] if "--out" in argv else "final/price-recommendations.xlsx"
    products = [json.load(open(os.path.join(LIVE, "products", f))) for f in os.listdir(os.path.join(LIVE, "products"))
                if f[:-5].isdigit()]
    root, path = tree()
    tr = truth()
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Recommendations"
    ws.append(["Item", "Name", "Brand id", "Animal", "Type", "Our cost", "Recommended price", "Margin %",
               "Method A (markup)", "Method B (unit price)", "Comparables used", "How", "Decision", "Approved price"])
    dv = DataValidation(type="list", formula1='"Approve,Use my price,Skip"', allow_blank=True)
    ws.add_data_validation(dv)
    ws2 = wb.create_sheet("Comparables")
    ws2.append(["Item", "Comparable product", "Variant", "SKU", "Price", "Cost", "Price/cost", "Price per g/ml/pc",
                "Same category", "Category", "Price source", "Link"])
    summary = []
    for it in items:
        comps = comparables(it, products, root, path, tr)
        rec = recommend(it, comps)
        margin = f"{(rec['recommended'] - it['cost']) / rec['recommended'] * 100:.1f}%" if rec["recommended"] else ""
        ws.append([it["key"], it["name"], it["brand_id"], it["animal"], it["type"], it["cost"], rec["recommended"], margin,
                   rec["method_a"], rec["method_b"], rec["n"], rec["note"], "", ""])
        dv.add(ws.cell(ws.max_row, 13))
        if rec["recommended"] is None or "BELOW" in rec["note"] or "LOW CONFIDENCE" in rec["note"] or "check" in rec["note"]:
            for c in range(1, 15):
                ws.cell(ws.max_row, c).fill = PatternFill("solid", fgColor="FBE3E1")
        for c in comps[:8]:
            ws2.append([it["key"], c["product"], c["variant"], c["sku"], c["price"], c["cost"], round(c["markup"], 2),
                        round(c["unit"], 3) if c["unit"] else None, "yes" if c["same_leaf"] else "", c["category"],
                        c["source"], c["link"]])
        summary.append((it["key"], it["cost"], rec["recommended"], rec["method_b"], rec["n"]))
    for w in (ws, ws2):
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "B2"
    for col, width in zip("ABCDEFGHIJKLMN", (12, 40, 9, 9, 12, 10, 14, 10, 14, 14, 12, 46, 14, 14)):
        ws.column_dimensions[col].width = width
    wb.save(os.path.join(ROOT, out))
    for s in summary:
        print(*s)


if __name__ == "__main__":
    main(sys.argv[1:])
