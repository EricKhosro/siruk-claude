#!/usr/bin/env python3
"""Printable lists of the per-kg scale codes (PLU) — one Excel for cats, one for dogs.

For the shop counter: the seller finds the 5-digit code to type on the scale.
Every ArmSoft good sold by weight (main unit 303 + Քաշային) is listed with its
scale code, ArmSoft name (what the till/scale shows), Siruk name + bag size and
the per-kg price from the ArmSoft price list; grouped under a big header per
brand. Species and brand come from the Siruk variant that carries the good's
code (accounting_code). Laid out for A4 landscape, fit to one page wide, header
repeated on every page.

    uv run --with openpyxl python scripts/scale-code-sheets.py \
        --goods Scale/armsoft-export-<date>.xml --prices Scale/armsoft-prices-<date>.xlsx \
        --out final/armsoft-<date>/
Needs the live dump in .siruk-cache/live/products/ (scripts/dump-products.py)
and .siruk-cache/brands.json + cats.json (GET /brands, GET /categories).
"""
import argparse, datetime, glob, importlib.util, json, os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
spec = importlib.util.spec_from_file_location('gaps', os.path.join(ROOT, 'scripts/armsoft-gaps.py'))
gaps = importlib.util.module_from_spec(spec); spec.loader.exec_module(gaps)

THIN = Side(style='thin', color='999999')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load(args):
    goods = {g['code']: g for g in gaps.parse_goods(args.goods)}
    prices = gaps.parse_prices(args.prices)
    cache = os.path.join(ROOT, '.siruk-cache')
    brands = {b['id']: b['name'] for b in json.load(open(os.path.join(cache, 'brands.json'), encoding='utf-8'))['data']}
    root = {}

    def walk(node, top):
        root[node['id']] = top
        for kid in node.get('children') or []:
            walk(kid, top)
    for node in json.load(open(os.path.join(cache, 'cats.json'), encoding='utf-8'))['data']:
        walk(node, node['name'])
    by_code = {}
    for f in glob.glob(os.path.join(cache, 'live/products/[0-9]*.json')):
        p = json.load(open(f, encoding='utf-8'))
        for v in p['variants']:
            if v.get('accounting_code'):
                by_code.setdefault(v['accounting_code'], []).append((p, v))
    rows = []
    for code, g in goods.items():
        if not (g['is_weight'] and g['unit'] == '303'):
            continue
        plu = (g['barcodes'].get('303') or [None])[0] or g['main_barcode']
        price = prices.get((code, '303'))
        hits = by_code.get(code, [])
        if not plu or not price or not hits:
            print(f'skipped {code} {g["name"]}: plu={plu} price={price} siruk={len(hits)}')
            continue
        p = hits[0][0]
        species = {root.get(c) for c in p['category_ids']} & {'Cat', 'Dog'}
        loose = [v for _, v in hits if (v['sku'] or '').endswith(('-KG', '-100G'))]
        per100 = next((v['price'] for v in loose if v['sku'].endswith('-100G')), None)
        bag_skus = {v['sku'].rsplit('-', 1)[0] for v in loose}
        bag = next((v for _, v in hits if v['sku'] in bag_skus), None)
        en = p['name'] + (f' — {bag["name"]}' if bag else '')
        for sp in species:
            rows.append(dict(species=sp, brand=brands.get(p['brand_id'], '?'), plu=plu, hy=g['name'],
                             en=en, price=price, per100=per100, code=code))
    return rows


def sheet(rows, species, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Cats' if species == 'Cat' else 'Dogs'
    widths = [12, 60, 64, 15]
    for col, w in zip('ABCD', widths):
        ws.column_dimensions[col].width = w
    title = 'ԿԱՏՈՒՆԵՐ — կշռվող կերեր, կշեռքի կոդեր' if species == 'Cat' else 'ՇՆԵՐ — կշռվող կերեր, կշեռքի կոդեր'
    ws.append([f'{title}   ({"Cats" if species == "Cat" else "Dogs"} — scale codes, per kg)'])
    ws.merge_cells('A1:D1')
    ws['A1'].font = Font(size=16, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26
    ws.append(['Կոդ (կշեռք)', 'Անվանում (ArmSoft)', 'Siruk', 'Գին / 1 կգ'])
    for c in ws[2]:
        c.font = Font(size=11, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='333333')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX
    ws.row_dimensions[2].height = 20
    ws.print_title_rows = '2:2'
    fills = ['B4C6E7', 'C6E0B4']
    for i, brand in enumerate(sorted({r['brand'] for r in rows}, key=str.lower)):
        items = sorted((r for r in rows if r['brand'] == brand), key=lambda r: r['plu'])
        ws.append([f'{brand}   ({len(items)})'])
        r0 = ws.max_row
        ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=4)
        cell = ws.cell(r0, 1)
        cell.font = Font(size=15, bold=True)
        cell.fill = PatternFill('solid', fgColor=fills[i % 2])
        cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        for c in range(1, 5):
            ws.cell(r0, c).border = Border(top=Side(style='thick'), bottom=Side(style='thick'))
        ws.row_dimensions[r0].height = 24
        for r in items:
            price = f'{r["price"]:,.0f} դր'.replace(',', ' ')
            if r['per100']:
                price += f'\n(100 գ — {r["per100"]:,.0f})'.replace(',', ' ')
            ws.append([r['plu'], r['hy'], r['en'], price])
            rr = ws.max_row
            ws.cell(rr, 1).font = Font(size=16, bold=True)
            ws.cell(rr, 1).number_format = '@'
            ws.cell(rr, 2).font = Font(size=11, bold=True)
            ws.cell(rr, 3).font = Font(size=10)
            ws.cell(rr, 4).font = Font(size=11, bold=True)
            for c in range(1, 5):
                ws.cell(rr, c).alignment = Alignment(vertical='center', wrap_text=True,
                                                     horizontal='center' if c in (1, 4) else 'left')
                ws.cell(rr, c).border = BOX
            ws.row_dimensions[rr].height = 30 if r['per100'] else 21
    ws.append([])
    ws.append([f'Siruk · {datetime.date.today():%d.%m.%Y} · {len(rows)} ապրանք'])
    ws.cell(ws.max_row, 1).font = Font(size=10, italic=True, color='666666')
    ps = ws.page_setup
    ps.orientation = 'landscape'
    ps.paperSize = ws.PAPERSIZE_A4
    ps.fitToWidth = 1
    ps.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = '&P / &N'
    ws.freeze_panes = 'A3'
    wb.save(path)
    print('wrote', path, len(rows), 'rows')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--goods', required=True)
    ap.add_argument('--prices', required=True)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    rows = load(a)
    os.makedirs(a.out, exist_ok=True)
    sheet([r for r in rows if r['species'] == 'Cat'], 'Cat', os.path.join(a.out, 'scale-codes-cats.xlsx'))
    sheet([r for r in rows if r['species'] == 'Dog'], 'Dog', os.path.join(a.out, 'scale-codes-dogs.xlsx'))


if __name__ == '__main__':
    main()
